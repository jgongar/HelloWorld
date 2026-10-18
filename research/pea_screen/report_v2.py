"""Build PEA screening v2 (report .md + .xlsx) from the saved scored universe. No network calls."""
import re, sys, unicodedata
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = "research"
DATE = "2026-10-18"
df = pd.read_csv(f"{ROOT}/data/{DATE}/scored_universe.csv")
df["passes_threshold"] = df["passes_threshold"].astype(str) == "True"
df["excluded"] = df["excluded"].fillna("").astype(str)
df["data_flag"] = df["data_flag"].fillna("").astype(str)
df["divergence"] = df["divergence"].fillna(0).astype(int)


def access_rank(f):
    f = str(f)
    if f.startswith("direct"): return 0
    if f.startswith("Xetra"): return 1
    if f.startswith("German secondary"): return 2
    return 3


def access_label(f):
    return ["Direct (Paris/Brussels/Amsterdam)", "Xetra/Frankfurt", "German secondary line (check liquidity)",
            "**Not online** (sale by phone only)"][access_rank(f)]


# Euronext Paris/Amsterdam/Brussels lines are buyable online even where the source tag says otherwise
EN = df.venue.fillna("").str.contains(r"Euronext (?:Growth |Access )?(?:Paris|Amsterdam|Brussels)", regex=True)
n_fixed = int((EN & ~df.fortuneo.fillna("").str.startswith("direct")).sum())
df.loc[EN, "fortuneo"] = "direct (Euronext P/A/B)"
df["acc"] = df["fortuneo"].map(access_rank)
clean = df[df.passes_threshold & (df.excluded == "") & (df.data_flag == "")].copy()
passing = df[df.passes_threshold & (df.excluded == "")].copy()


def tier(r):
    rv = 0.5 if pd.isna(r.p_revisions) else r.p_revisions
    q = 0.5 if pd.isna(r.p_quality) else r.p_quality
    if rv >= 0.6 and q >= 0.5: return "A"
    if rv < 0.4 or q < 0.35: return "C"
    return "B"


clean["tier"] = clean.apply(tier, axis=1)
clean["high_risk"] = (clean["size"] == "S") & (clean.upside > 1.5) & (clean.n_analysts < 3)
# Access preference: small bonus for easier Fortuneo access, then rank by score
clean["sel_score"] = clean.score + clean.acc.map({0: 3, 1: 2, 2: 0, 3: 0})
pool = clean.sort_values("sel_score", ascending=False)
picked, n_off = [], 0
for i, r in pool.iterrows():
    if r.acc == 3:
        if n_off >= 10: continue
        n_off += 1
    picked.append(i)
    if len(picked) == 60: break
leads = clean.loc[picked].sort_values("score", ascending=False).copy()
leads["rank"] = range(1, len(leads) + 1)


# ---------- formatting helpers ----------
def nz(x): return x is not None and not (isinstance(x, float) and np.isnan(x))
def pct(x, d=0, sign=True):
    if not nz(x): return "n/a"
    return f"{x*100:+.{d}f}%" if sign else f"{x*100:.{d}f}%"
def num(x, d=1, suf=""):
    return f"{x:.{d}f}{suf}" if nz(x) else "n/a"
def money(x):
    if not nz(x): return "n/a"
    return f"€{x/1e9:.1f}bn" if x >= 1e9 else f"€{x/1e6:.0f}m" if x >= 1e7 else f"€{x/1e6:.1f}m" if x >= 1e6 else f"€{x/1e3:.0f}k"
def px(x, c):
    if not nz(x): return "n/a"
    return f"{c} {x:,.2f}" if x < 1000 else f"{c} {x:,.0f}"
def cnt(x): return int(x) if nz(x) else 0


def what(summary):
    s = str(summary) if nz(summary) else ""
    sents = re.split(r"(?<=[a-z\)])\.\s+(?=[A-Z])", s)
    out = ". ".join(sents[:2]).strip()
    if len(out) > 330: out = out[:327].rsplit(" ", 1)[0] + "…"
    return out.rstrip(".") + "." if out else "No description available."


NICHE_IND = ("Semiconductor", "Specialty Chemicals", "Chemicals", "Aerospace & Defense", "Electronic Components",
             "Scientific & Technical", "Specialty Industrial Machinery", "Medical Instruments", "Medical Devices",
             "Diagnostics", "Electrical Equipment", "Industrial Distribution", "Pollution", "Metal Fabrication",
             "Communication Equipment", "Computer Hardware", "Biotechnology", "Drug Manufacturers - Specialty",
             "Solar", "Uranium", "Other Industrial Metals", "Building Products", "Tools & Accessories",
             "Software - Application", "Software - Infrastructure", "Information Technology Services", "Engineering")
THEME_KW = {"AI / data centres": r"data cent|semiconductor|chip|wafer|AI\b|artificial intelligence|photonic|laser",
            "defence": r"defen[cs]e|military|ammunition|armou?r",
            "chemicals / materials": r"chemical|polymer|coating|pigment|resin|catalyst|lithium|graphite|specialty material",
            "energy transition": r"wind|solar|hydrogen|battery|electrification|grid|renewable|heat pump",
            "healthcare niche": r"medical|diagnostic|surgical|pharma|implant|dental|radio",
            "automation / industrial tech": r"automation|robot|sensor|machine vision|measurement|instrument"}


def themes(r):
    txt = f"{r.industry} {r.summary}"
    return [k for k, v in THEME_KW.items() if re.search(v, txt, re.I)]


def lead_types(r):
    t = []
    if (nz(r.p_revisions) and r.p_revisions >= 0.65) or (nz(r.rev90_fy1) and r.rev90_fy1 > 0.03): t.append("Revisions↑")
    if r.divergence == 1: t.append("Forecasts up / price down")
    if nz(r.p_valuation) and r.p_valuation >= 0.65: t.append("Cheap vs sector")
    if (nz(r.net_cash_to_mcap) and r.net_cash_to_mcap >= 0.15) or (nz(r.p_book) and 0 < r.p_book < 1):
        t.append("Hidden assets / net cash")
    if any(k in str(r.industry) for k in NICHE_IND) and (nz(r.p_quality) and r.p_quality >= 0.55 or themes(r)):
        t.append("Niche leader / supply-chain theme")
    return t or ["Upside to consensus"]


def risks(r):
    k = []
    if nz(r.net_debt_ebitda) and r.net_debt_ebitda > 3: k.append(f"high leverage (net debt {r.net_debt_ebitda:.1f}× EBITDA)")
    if (nz(r.rev90_fy1) and r.rev90_fy1 < -0.03) or (cnt(r.down30_fy1) > cnt(r.up30_fy1) and nz(r.rev30_fy1) and r.rev30_fy1 < 0):
        k.append("estimate cuts")
    if r.upside > 1.0: k.append(f"extreme upside ({pct(r.upside)}) — check the targets are fresh")
    if r.n_analysts <= 2: k.append(f"only {int(r.n_analysts)} analysts")
    if nz(r.fcf_conversion) and r.fcf_conversion < 0.5: k.append(f"weak cash conversion ({r.fcf_conversion:.0%} FCF/NI)")
    if nz(r.op_margin) and r.op_margin < 0: k.append("loss-making at operating level")
    if nz(r.target_dispersion) and r.target_dispersion > 0.5: k.append("wide spread of analyst targets")
    if r.acc == 2: k.append("German secondary line: spreads may be wide")
    if r.acc == 3: k.append("not buyable online at Fortuneo")
    if r.high_risk: k.append("**HIGH RISK: small cap, >150% upside, <3 analysts**")
    return k or ["nothing obvious in the data"]


def rationale(r, types):
    bits = []
    if nz(r.rev90_fy1): bits.append(f"next-year EPS forecast {pct(r.rev90_fy1,1)} over 90 days ({cnt(r.up30_fy1)} up / {cnt(r.down30_fy1)} down in 30 days)")
    if nz(r.eps_growth_fy1): bits.append(f"consensus EPS growth {pct(r.eps_growth_fy1)} next year")
    if nz(r.fwd_pe): bits.append(f"{r.fwd_pe:.1f}× forward earnings")
    if nz(r.roic): bits.append(f"ROIC {pct(r.roic,0,False)}")
    th = themes(r)
    s = f"{pct(r.upside)} to the mean target of {int(r.n_analysts)} analysts; " + ", ".join(bits) + "."
    if th: s += f" Theme: {', '.join(th)}."
    if r.divergence == 1: s += " Forecasts rose while the share price fell — the market may not have noticed."
    return s


# ---------- run-1 comparison ----------
def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9 ]", " ", s).split()


run1 = []
for line in open(f"{ROOT}/pea-screening.md", encoding="utf-8"):
    m = re.match(r"\| (\d+) \| \*\*(.+?)\*\* \|", line)
    if m and int(m.group(1)) <= 48: run1.append(m.group(2))
ALIAS = {"AT&S": ["at", "s"], "DO & CO": ["do", "co"], "ALK-Abelló": ["alk"]}
cmp_rows = []
lead_idx = set(leads.index)
for nm in run1:
    key = ALIAS.get(nm, norm(nm))
    cand = df[df.name.map(lambda x: norm(x)[:len(key)] == key)]
    if cand.empty and len(key) > 1:
        cand = df[df.name.map(lambda x: norm(x)[:1] == key[:1])]
    if cand.empty:
        cmp_rows.append((nm, "—", "not matched in the v2 data (name or symbol missing)")); continue
    r = cand.sort_values("mcap_eur", ascending=False).iloc[0]
    if r.name in lead_idx:
        st = f"**still in** — #{int(leads.loc[r.name,'rank'])}, tier {leads.loc[r.name,'tier']}"
    elif r.excluded:
        st = f"dropped — excluded ({r.excluded})"
    elif r.data_flag:
        st = f"dropped — data flag ({r.data_flag})"
    elif nz(r.n_analysts) and r.n_analysts < 2 or not nz(r.upside):
        st = "dropped — fewer than 2 analysts in Yahoo data"
    elif not r.passes_threshold:
        st = f"dropped — threshold: upside {pct(r.upside)} vs {pct(r.min_upside,0)} needed"
    else:
        st = f"passes threshold but outside top 60 (score {r.score:.0f})"
    cmp_rows.append((nm, f"{r['name']} ({r.symbol})", st))
n_still = sum("still in" in c[2] for c in cmp_rows)

# ---------- report ----------
L = []
w = L.append
w(f"# PEA screening v2: leads where consensus is well above the price\n")
w(f"*Run 2 · data from Yahoo Finance, fetched {DATE} · built from `research/data/{DATE}/scored_universe.csv` · not investment advice — verify before acting.*\n")
w("## 1. Summary\n")
acc_counts = leads.acc.value_counts().reindex([0, 1, 2, 3]).fillna(0).astype(int).tolist()
w(f"- Universe: **{len(df):,} PEA-eligible companies** (EU/EEA registered office, market cap ≥ €50m, ≥ €100k traded per day). "
  f"{(df.excluded!='').sum()} excluded (banks, insurers, REITs, non-EU domicile, pre-revenue biotech).")
w(f"- {len(passing)} pass the upside threshold (≥ 2 analysts; upside ≥ 35% small, ≥ 25% mid, ≥ 15% large); "
  f"{len(clean)} of those have clean data (no `data_flag`).")
w(f"- **Main list: 60 leads**, ranked by score (revisions 25, upside 20, quality 20, growth 15, valuation vs sector 15, balance sheet 5). "
  f"Fortuneo access: {acc_counts[0]} direct, {acc_counts[1]} Xetra/Frankfurt, {acc_counts[2]} German secondary line, {acc_counts[3]} not online (flagged).")
tc = leads.tier.value_counts()
w(f"- Tiers: **A** {tc.get('A',0)} (revisions pillar ≥ 0.6 and quality ≥ 0.5), **B** {tc.get('B',0)}, **C** {tc.get('C',0)} (revisions < 0.4 or quality < 0.35).")
w(f"- {int(leads.high_risk.sum())} leads carry the high-risk flag (small cap, upside > 150%, fewer than 3 analysts).\n")

w("## 2. Main list: 60 leads\n")
w("| # | Company | Country / venue | Size | Fortuneo | Price → target (n) | Upside | Rev. 90d FY+1 | Fwd P/E | ROIC | Score | Tier | Lead type |")
w("|---|---|---|---|---|---|---|---|---|---|---|---|---|")
short_acc = {0: "Direct", 1: "Xetra", 2: "DE 2nd line", 3: "**Not online**"}
for _, r in leads.iterrows():
    flag = " ⚠" if r.high_risk else ""
    w(f"| {r['rank']} | **{r['name']}**{flag} | {r.country} / {r.venue} | {r['size']} | {short_acc[r.acc]} | "
      f"{px(r.price,r.currency)} → {px(r.target_mean,r.currency)} ({int(r.n_analysts)}) | {pct(r.upside)} | {pct(r.rev90_fy1,1)} | "
      f"{num(r.fwd_pe)} | {pct(r.roic,0,False)} | {r.score:.0f} | {r.tier} | {', '.join(lead_types(r))} |")
w("\n⚠ = high risk (small cap, > 150% upside, fewer than 3 analysts).\n")

w("## 3. Lead cards\n")
for t in "ABC":
    sub = leads[leads.tier == t]
    if sub.empty: continue
    w(f"### Tier {t}\n")
    for _, r in sub.iterrows():
        ty = lead_types(r)
        nc = r.net_cash_to_mcap
        bs = (f"net cash {pct(nc,0,False)} of market cap" if nz(nc) and nc > 0 else
              f"net debt {num(r.net_debt_ebitda)}× EBITDA" if nz(r.net_debt_ebitda) else "n/a")
        w(f"#### {r['rank']}. {r['name']} ({r.symbol}) — tier {r.tier}, score {r.score:.0f}\n")
        w(f"- **What it does:** {what(r.summary)}")
        w(f"- **Listing:** {r.country}, {r.venue} · {r.industry} · market cap {money(r.mcap_eur)}, {money(r.adv_eur)}/day traded · ISIN {r['isin'] if nz(r['isin']) else 'n/a'}")
        w(f"- **Fortuneo:** {access_label(r.fortuneo)}")
        w(f"- **Consensus:** {px(r.price,r.currency)} → {px(r.target_mean,r.currency)} mean (range {px(r.target_low,r.currency)}–{px(r.target_high,r.currency)}, {int(r.n_analysts)} analysts), "
          f"upside **{pct(r.upside)}**; recommendation {num(r.reco_mean)} (1 = strong buy)")
        w(f"- **Revisions:** FY+1 EPS {pct(r.rev30_fy1,1)} (30d) / {pct(r.rev90_fy1,1)} (90d), {cnt(r.up30_fy1)} up / {cnt(r.down30_fy1)} down; "
          f"current year {pct(r.rev30_fy0,1)} / {pct(r.rev90_fy0,1)}, {cnt(r.up30_fy0)} up / {cnt(r.down30_fy0)} down")
        w(f"- **Valuation:** fwd P/E {num(r.fwd_pe)}, EV/EBITDA {num(r.ev_ebitda)}, FCF yield {pct(r.fcf_yield,1,False)}, P/B {num(r.p_book)}")
        w(f"- **Quality:** ROIC {pct(r.roic,0,False)}, operating margin {pct(r.op_margin,0,False)}, Piotroski {int(r.piotroski7) if nz(r.piotroski7) else 'n/a'}/7, FCF conversion {pct(r.fcf_conversion,0,False)}")
        w(f"- **Growth:** sales {pct(r.rev_growth_fy1)} and EPS {pct(r.eps_growth_fy1)} next year (consensus); 3-year sales CAGR {pct(r.rev_cagr_3y)}")
        w(f"- **Balance sheet:** {bs}")
        w(f"- **Momentum:** {pct(r.ret_13w)} (13w), {pct(r.ret_52w)} (52w), {pct(r.drawdown_52w)} from 52-week high; divergence flag: {'**yes**' if r.divergence else 'no'}")
        w(f"- **Rationale [{' · '.join(ty)}]:** {rationale(r, ty)}")
        w(f"- **Key risk:** {'; '.join(risks(r))}\n")

w("## 4. Uncovered gems (top 15)\n")
w("No analyst consensus (so no upside figure), but high score on quality, growth, valuation and balance sheet. Fortuneo-accessible only.\n")
gems = df[(df.coverage == "uncovered") & (df.acc < 3) & (df.excluded == "") & (df.data_flag == "")].nlargest(15, "score")
w("| # | Company | Country / venue | Fortuneo | Industry | Mkt cap | Fwd/trailing P/E | EV/EBITDA | ROIC | Op. margin | Net cash / mcap | Score |")
w("|---|---|---|---|---|---|---|---|---|---|---|---|")
for k, (_, r) in enumerate(gems.iterrows(), 1):
    pe = r.fwd_pe if nz(r.fwd_pe) else r.trailing_pe
    w(f"| {k} | **{r['name']}** | {r.country} / {r.venue} | {short_acc[r.acc]} | {r.industry} | {money(r.mcap_eur)} | {num(pe)} | {num(r.ev_ebitda)} | "
      f"{pct(r.roic,0,False)} | {pct(r.op_margin,0,False)} | {pct(r.net_cash_to_mcap,0)} | {r.score:.0f} |")

w("\n## 5. Divergence: forecasts up, price down (top 15)\n")
w("Next-year EPS forecast up > 2% over 90 days while the share price fell > 5% over 13 weeks; passes the upside threshold.\n")
div = clean[clean.divergence == 1].nlargest(15, "score")
w("| # | Company | Country | Fortuneo | Upside | Rev. 90d FY+1 | Price 13w | Fwd P/E | Score | In main list |")
w("|---|---|---|---|---|---|---|---|---|---|")
for k, (_, r) in enumerate(div.iterrows(), 1):
    inl = f"#{int(leads.loc[r.name,'rank'])}" if r.name in lead_idx else "—"
    w(f"| {k} | **{r['name']}** | {r.country} | {short_acc[r.acc]} | {pct(r.upside)} | {pct(r.rev90_fy1,1)} | {pct(r.ret_13w)} | {num(r.fwd_pe)} | {r.score:.0f} | {inl} |")

w(f"\n## 6. Comparison with run 1\n")
w(f"Run 1 (`research/pea-screening.md`) had 48 core leads built from web research. **{n_still} are still in the v2 main list.** "
  "Run 1 used fresh broker targets from news; v2 uses only the Yahoo consensus as of 18 Oct, so differences in upside are expected.\n")
w("| Run-1 lead | Matched in v2 data | Status in v2 |")
w("|---|---|---|")
for a, b, c in cmp_rows: w(f"| {a} | {b} | {c} |")

w("\n## 7. Method and caveats\n")
w("- **Source:** all figures come from Yahoo Finance, fetched on 18 Oct 2026. For small caps the consensus may be one or two brokers, sometimes stale or in the wrong currency; `data_flag` rows (upside outside −80%/+300%) are left out of the lists.")
w("- **Scoring:** each pillar is the average of percentile ranks across the universe (0–1); score = 25 revisions + 20 upside + 20 quality + 15 growth + 15 valuation vs sector + 5 balance sheet. ROIC uses EBIT × 0.75.")
w("- **Selection:** clean passing names sorted by score with a small bonus for easier Fortuneo access (direct +3, Xetra +2); at most 10 `not online` names.")
w("- **Lead types** and themes are inferred automatically from industry, the Yahoo business description and the numbers; check them.")
w("- **Coverage gaps:** exchanges in Croatia, Slovenia, Bulgaria, Slovakia, Cyprus and Malta are not covered; some German symbols are missing (e.g. K+S).")
w(f"- **Fortuneo access** follows the official tariff brochure (https://www.fortuneo.fr/files/tarifs_fortuneo.pdf): Paris/Brussels/Amsterdam online; Xetra/Frankfurt at 0.20% (min €20, +€30 foreign fee, €400 minimum order); Milan, Madrid, Nordics, Vienna, Lisbon sale-only by phone.")
w("- **Not investment advice — verify every figure (price, targets, latest results) before acting.**")
open(f"{ROOT}/pea-screening-v2.md", "w", encoding="utf-8").write("\n".join(L) + "\n")

# ---------- Excel ----------
HDR = {"symbol": "Symbol", "name": "Company", "isin": "ISIN", "venue": "Venue", "fortuneo": "Fortuneo access", "all_lines": "All listings",
       "country": "Country", "sector": "Sector", "industry": "Industry", "summary": "Business description", "website": "Website",
       "currency": "Currency", "price": "Price", "mcap_eur": "Market cap (€)", "adv_eur": "Avg daily traded (€)",
       "n_analysts": "Analysts", "target_mean": "Target mean", "target_median": "Target median", "target_high": "Target high",
       "target_low": "Target low", "reco_mean": "Recommendation (1=strong buy)", "upside": "Upside to mean target",
       "target_dispersion": "Target dispersion", "rev30_fy0": "EPS rev 30d (FY0)", "rev90_fy0": "EPS rev 90d (FY0)",
       "rev30_fy1": "EPS rev 30d (FY+1)", "rev90_fy1": "EPS rev 90d (FY+1)", "up30_fy1": "Analysts up 30d (FY+1)",
       "down30_fy1": "Analysts down 30d (FY+1)", "eps_growth_fy1": "EPS growth FY+1", "rev_growth_fy1": "Sales growth FY+1",
       "rev_cagr_3y": "Sales CAGR 3y", "fwd_pe": "Forward P/E", "trailing_pe": "Trailing P/E", "ev_ebitda": "EV/EBITDA",
       "ev_sales": "EV/Sales", "p_book": "P/Book", "div_yield": "Dividend yield", "fcf_yield": "FCF yield", "peg": "PEG",
       "net_cash_to_mcap": "Net cash / mcap", "net_debt_ebitda": "Net debt / EBITDA", "op_margin": "Operating margin",
       "gross_margin": "Gross margin", "roe": "ROE", "roic": "ROIC", "op_margin_std": "Op. margin volatility",
       "fcf_conversion": "FCF conversion", "piotroski7": "Piotroski (0-7)", "ret_13w": "Return 13w", "ret_26w": "Return 26w",
       "ret_52w": "Return 52w", "drawdown_52w": "From 52w high", "divergence": "Divergence flag", "excluded": "Excluded reason",
       "p_revisions": "Pillar: revisions", "p_upside": "Pillar: upside", "p_quality": "Pillar: quality", "p_growth": "Pillar: growth",
       "p_valuation": "Pillar: valuation", "p_balance": "Pillar: balance sheet", "score": "Score (0-100)", "coverage": "Coverage",
       "size": "Size", "min_upside": "Min upside required", "passes_threshold": "Passes threshold", "data_flag": "Data flag",
       "rank": "Rank", "tier": "Tier", "high_risk": "High risk flag", "lead_type": "Lead type", "key_risk": "Key risk"}
PCT = {"upside", "rev30_fy0", "rev90_fy0", "rev30_fy1", "rev90_fy1", "eps_growth_fy1", "rev_growth_fy1", "rev_cagr_3y", "div_yield",
       "fcf_yield", "net_cash_to_mcap", "op_margin", "gross_margin", "roe", "roic", "fcf_conversion", "ret_13w", "ret_26w",
       "ret_52w", "drawdown_52w", "min_upside", "eps_growth_fy0", "rev_growth_fy0"}
base_cols = [c for c in df.columns if c not in ("acc",) and not c.startswith("r_")]
leads["lead_type"] = leads.apply(lambda r: ", ".join(lead_types(r)), axis=1)
leads["key_risk"] = leads.apply(lambda r: "; ".join(risks(r)).replace("**", ""), axis=1)
lead_cols = ["rank", "tier", "symbol", "name", "country", "venue", "fortuneo", "industry", "size", "mcap_eur", "adv_eur", "currency",
             "price", "target_mean", "n_analysts", "upside", "rev30_fy1", "rev90_fy1", "up30_fy1", "down30_fy1", "fwd_pe", "ev_ebitda",
             "fcf_yield", "roic", "op_margin", "piotroski7", "eps_growth_fy1", "rev_growth_fy1", "net_cash_to_mcap", "net_debt_ebitda",
             "ret_13w", "ret_52w", "divergence", "score", "high_risk", "lead_type", "key_risk", "isin", "website", "summary"]

wb = Workbook(); wb.remove(wb.active)
for title, frame, cols in [("Leads", leads, lead_cols), ("All passing", passing.sort_values("score", ascending=False), base_cols),
                           ("Universe", df.sort_values("score", ascending=False), base_cols)]:
    ws = wb.create_sheet(title)
    ws.append([HDR.get(c, c.replace("_", " ").capitalize()) for c in cols])
    for row in frame[cols].itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, "item") else v) for v in row])
    for j, c in enumerate(cols, 1):
        cell = ws.cell(1, j); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = 45 if c in ("summary", "key_risk", "all_lines") else 28 if c in ("name", "industry", "fortuneo", "lead_type") else 13
        fmt = "0.0%" if c in PCT else "#,##0" if c in ("mcap_eur", "adv_eur") else "0.00" if frame[c].dtype.kind == "f" else None
        if fmt:
            for k in range(2, ws.max_row + 1): ws.cell(k, j).number_format = fmt
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
wb.save(f"{ROOT}/pea-screening-v2.xlsx")

print(f"universe {len(df)} passing {len(passing)} clean {len(clean)} leads {len(leads)} tiers {dict(tc)} access {acc_counts} still_in {n_still}/{len(run1)}")
for _, r in leads.head(10).iterrows():
    print(f"{r['rank']}|{r['name']}|{r.country}|{r.venue}|{short_acc[r.acc]}|{r.industry}|up {pct(r.upside)}|n{int(r.n_analysts)}|rev90 {pct(r.rev90_fy1,1)}|pe {num(r.fwd_pe)}|roic {pct(r.roic,0,False)}|{r.tier}|{r.score:.0f}|{lead_types(r)}|{what(r.summary)[:120]}")
